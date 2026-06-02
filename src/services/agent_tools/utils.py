"""本地实用工具：无需外部密钥，开箱即用。

参照 hermes 的工具粒度，提供日期时间、计算、JSON/正则提取、文本统计、
文件下载/保存/列举、文档抽取等常用能力。所有工具均为 LangChain
``StructuredTool``，由 :func:`build_util_tools` 工厂注入租户上下文后返回。
"""

from __future__ import annotations

import ast
import json
import logging
import math
import operator
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, List, Optional

from langchain_core.tools import StructuredTool
from pydantic import BaseModel, Field

logger = logging.getLogger(__name__)

# 单个工具返回给模型的文本上限，避免把大文件/长网页整段灌进上下文。
_MAX_CHARS = 12000
# 下载文件大小上限（字节）。
_MAX_DOWNLOAD_BYTES = 50 * 1024 * 1024


def _err(message: str, **extra: Any) -> str:
    payload = {"error": str(message)}
    payload.update(extra)
    return json.dumps(payload, ensure_ascii=False)


def _ok(data: Any) -> str:
    return json.dumps(data, ensure_ascii=False, default=str)


def _truncate(text: str) -> str:
    if len(text) > _MAX_CHARS:
        return text[:_MAX_CHARS] + f"\n…（已截断，共 {len(text)} 字符）"
    return text


# ---------------------------------------------------------------------------
# 安全算术表达式求值
# ---------------------------------------------------------------------------

_ALLOWED_BINOPS = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
    ast.FloorDiv: operator.floordiv,
    ast.Mod: operator.mod,
    ast.Pow: operator.pow,
}
_ALLOWED_UNARYOPS = {ast.UAdd: operator.pos, ast.USub: operator.neg}
_ALLOWED_FUNCS = {
    "sqrt": math.sqrt, "log": math.log, "log10": math.log10, "exp": math.exp,
    "sin": math.sin, "cos": math.cos, "tan": math.tan,
    "floor": math.floor, "ceil": math.ceil, "abs": abs, "round": round,
    "pow": math.pow, "factorial": math.factorial,
}
_ALLOWED_NAMES = {"pi": math.pi, "e": math.e, "tau": math.tau}


def _safe_eval(node: ast.AST) -> float:
    if isinstance(node, ast.Constant):
        if isinstance(node.value, (int, float)):
            return node.value
        raise ValueError("不支持的常量类型")
    if isinstance(node, ast.BinOp) and type(node.op) in _ALLOWED_BINOPS:
        return _ALLOWED_BINOPS[type(node.op)](_safe_eval(node.left), _safe_eval(node.right))
    if isinstance(node, ast.UnaryOp) and type(node.op) in _ALLOWED_UNARYOPS:
        return _ALLOWED_UNARYOPS[type(node.op)](_safe_eval(node.operand))
    if isinstance(node, ast.Name) and node.id in _ALLOWED_NAMES:
        return _ALLOWED_NAMES[node.id]
    if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
        fn = _ALLOWED_FUNCS.get(node.func.id)
        if fn is None:
            raise ValueError(f"不支持的函数: {node.func.id}")
        return fn(*[_safe_eval(a) for a in node.args])
    raise ValueError("表达式包含不支持的语法")


# ---------------------------------------------------------------------------
# 工具实现（每个返回 JSON 字符串）
# ---------------------------------------------------------------------------


class _DatetimeArgs(BaseModel):
    tz_offset_hours: Optional[float] = Field(
        default=None, description="相对 UTC 的时区偏移小时数，如东八区填 8。留空则用本机时区。"
    )


def _get_datetime(tz_offset_hours: Optional[float] = None) -> str:
    if tz_offset_hours is None:
        now = datetime.now().astimezone()
    else:
        from datetime import timedelta
        now = datetime.now(timezone(timedelta(hours=tz_offset_hours)))
    return _ok({
        "iso": now.isoformat(),
        "date": now.strftime("%Y-%m-%d"),
        "time": now.strftime("%H:%M:%S"),
        "weekday": now.strftime("%A"),
        "tz": now.strftime("%z") or "local",
    })


class _CalcArgs(BaseModel):
    expression: str = Field(description="算术表达式，如 '(3+4)*2' 或 'sqrt(16)+log(e)'。")


def _calculator(expression: str) -> str:
    try:
        tree = ast.parse(expression, mode="eval")
        result = _safe_eval(tree.body)
        return _ok({"expression": expression, "result": result})
    except Exception as e:
        return _err(f"无法计算表达式: {e}", expression=expression)


class _JsonQueryArgs(BaseModel):
    data: str = Field(description="JSON 文本（对象或数组）。")
    path: str = Field(
        description="点/方括号路径，如 'a.b[0].c'。留空或 '.' 返回顶层。",
    )


def _json_query(data: str, path: str = "") -> str:
    try:
        obj = json.loads(data)
    except json.JSONDecodeError as e:
        return _err(f"输入不是合法 JSON: {e}")
    tokens = re.findall(r"[^.\[\]]+|\[\d+\]", path or "")
    cur = obj
    try:
        for tok in tokens:
            if tok.startswith("[") and tok.endswith("]"):
                cur = cur[int(tok[1:-1])]
            else:
                cur = cur[tok]
    except (KeyError, IndexError, TypeError) as e:
        return _err(f"路径 '{path}' 无法解析: {e}")
    return _ok({"path": path, "value": cur})


class _RegexArgs(BaseModel):
    text: str = Field(description="待搜索的原始文本。")
    pattern: str = Field(description="Python 正则表达式。")
    group: int = Field(default=0, description="返回的捕获组序号，0 表示整体匹配。")
    max_matches: int = Field(default=50, description="最多返回的匹配数量。")


def _regex_extract(text: str, pattern: str, group: int = 0, max_matches: int = 50) -> str:
    try:
        rx = re.compile(pattern)
    except re.error as e:
        return _err(f"正则表达式无效: {e}")
    matches: List[str] = []
    for m in rx.finditer(text):
        try:
            matches.append(m.group(group))
        except IndexError:
            return _err(f"捕获组 {group} 不存在")
        if len(matches) >= max_matches:
            break
    return _ok({"count": len(matches), "matches": matches})


class _TextStatsArgs(BaseModel):
    text: str = Field(description="待统计的文本。")


def _text_stats(text: str) -> str:
    chinese = len(re.findall(r"[一-鿿]", text))
    words = len(re.findall(r"\b\w+\b", text))
    return _ok({
        "chars": len(text),
        "chars_no_space": len(re.sub(r"\s", "", text)),
        "chinese_chars": chinese,
        "words": words,
        "lines": text.count("\n") + 1 if text else 0,
    })


def build_util_tools(settings: Any, tenant_uuid: Optional[str]) -> List[StructuredTool]:
    """构建本地实用工具，闭包捕获 settings/tenant_uuid。"""

    def _output_dir() -> Path:
        return Path(settings.get_tenant_output_dir(tenant_uuid))

    def _resolve_under_output(name: str) -> Path:
        """把文件名/相对路径限制在租户产物目录内，防止越权写。"""
        base = _output_dir().resolve()
        candidate = (base / name).resolve()
        if base not in candidate.parents and candidate != base:
            raise ValueError("禁止写入租户产物目录之外的位置")
        return candidate

    # ---- download_file -----------------------------------------------------
    class _DownloadArgs(BaseModel):
        url: str = Field(description="要下载的 http/https 直链。")
        filename: str = Field(description="保存到租户产物目录下的文件名（不含路径）。")

    def _download_file(url: str, filename: str) -> str:
        if not url.lower().startswith(("http://", "https://")):
            return _err("仅支持 http/https 链接")
        try:
            dest = _resolve_under_output(Path(filename).name)
        except ValueError as e:
            return _err(str(e))
        try:
            import requests
            with requests.get(url, stream=True, timeout=60) as r:
                r.raise_for_status()
                total = 0
                with open(dest, "wb") as f:
                    for chunk in r.iter_content(chunk_size=65536):
                        total += len(chunk)
                        if total > _MAX_DOWNLOAD_BYTES:
                            f.close()
                            dest.unlink(missing_ok=True)
                            return _err("文件超过 50MB 上限，已取消")
                        f.write(chunk)
            return _ok({"saved_to": str(dest), "bytes": total})
        except Exception as e:
            return _err(f"下载失败: {e}")

    # ---- list_outputs ------------------------------------------------------
    class _ListOutputsArgs(BaseModel):
        pass

    def _list_outputs() -> str:
        base = _output_dir()
        files = []
        for p in sorted(base.rglob("*")):
            if p.is_file():
                files.append({
                    "name": str(p.relative_to(base)),
                    "bytes": p.stat().st_size,
                })
        return _ok({"dir": str(base), "count": len(files), "files": files})

    # ---- save_text_file ----------------------------------------------------
    class _SaveTextArgs(BaseModel):
        filename: str = Field(description="保存到租户产物目录下的文件名（如 '作文.txt'）。")
        content: str = Field(description="要写入的文本内容。")

    def _save_text_file(filename: str, content: str) -> str:
        try:
            dest = _resolve_under_output(Path(filename).name)
        except ValueError as e:
            return _err(str(e))
        try:
            dest.write_text(content, encoding="utf-8")
            return _ok({"saved_to": str(dest), "bytes": len(content.encode("utf-8"))})
        except Exception as e:
            return _err(f"保存失败: {e}")

    # ---- read_document -----------------------------------------------------
    class _ReadDocArgs(BaseModel):
        path: str = Field(
            description="文档路径；相对路径按租户产物目录解析。支持 pdf/docx/xlsx/txt/md/csv/json。"
        )

    def _read_document(path: str) -> str:
        p = Path(path)
        if not p.is_absolute():
            p = _output_dir() / path
        if not p.exists():
            return _err(f"文件不存在: {p}")
        ext = p.suffix.lower()
        try:
            if ext == ".pdf":
                text = _extract_pdf(p)
            elif ext == ".docx":
                import docx
                text = "\n".join(par.text for par in docx.Document(str(p)).paragraphs)
            elif ext in (".xlsx", ".xlsm"):
                import openpyxl
                wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
                rows = []
                for ws in wb.worksheets:
                    rows.append(f"# Sheet: {ws.title}")
                    for row in ws.iter_rows(values_only=True):
                        rows.append("\t".join("" if c is None else str(c) for c in row))
                text = "\n".join(rows)
            elif ext in (".txt", ".md", ".csv", ".json", ".log", ".yaml", ".yml"):
                text = p.read_text(encoding="utf-8", errors="replace")
            else:
                return _err(f"不支持的文档类型: {ext}")
        except Exception as e:
            return _err(f"读取文档失败: {e}")
        return _ok({"path": str(p), "type": ext, "text": _truncate(text)})

    def _extract_pdf(p: Path) -> str:
        try:
            import pdfplumber
            with pdfplumber.open(str(p)) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        except Exception:
            from pypdf import PdfReader
            reader = PdfReader(str(p))
            return "\n".join(page.extract_text() or "" for page in reader.pages)

    return [
        StructuredTool.from_function(
            func=_get_datetime, name="get_datetime",
            description="获取当前日期、时间与星期，可指定时区偏移。", args_schema=_DatetimeArgs,
        ),
        StructuredTool.from_function(
            func=_calculator, name="calculator",
            description="安全计算算术表达式，支持 + - * / // % ** 与 sqrt/log/sin 等常用函数。",
            args_schema=_CalcArgs,
        ),
        StructuredTool.from_function(
            func=_json_query, name="json_query",
            description="按点/方括号路径从 JSON 文本中取值，如 path='items[0].name'。",
            args_schema=_JsonQueryArgs,
        ),
        StructuredTool.from_function(
            func=_regex_extract, name="regex_extract",
            description="用正则从文本中提取所有匹配项，可指定捕获组。", args_schema=_RegexArgs,
        ),
        StructuredTool.from_function(
            func=_text_stats, name="text_stats",
            description="统计文本的字符数、中文字数、词数与行数。", args_schema=_TextStatsArgs,
        ),
        StructuredTool.from_function(
            func=_download_file, name="download_file",
            description="下载 http/https 直链到当前租户的产物目录（上限 50MB）。",
            args_schema=_DownloadArgs,
        ),
        StructuredTool.from_function(
            func=_list_outputs, name="list_outputs",
            description="列出当前租户产物目录下的所有文件及大小。", args_schema=_ListOutputsArgs,
        ),
        StructuredTool.from_function(
            func=_save_text_file, name="save_text_file",
            description="把文本保存到当前租户的产物目录（前端下载列表可见）。",
            args_schema=_SaveTextArgs,
        ),
        StructuredTool.from_function(
            func=_read_document, name="read_document",
            description="抽取 pdf/docx/xlsx/txt/md/csv/json 文档的纯文本内容。",
            args_schema=_ReadDocArgs,
        ),
    ]
